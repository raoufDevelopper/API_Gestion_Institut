from decimal import Decimal

from django.db.models import Avg, Count, Q, Sum

from django.utils import timezone

from rest_framework.decorators import api_view, permission_classes

from rest_framework.permissions import IsAuthenticated

from rest_framework.response import Response

from apps.authentification.decorators import permission_requise

from apps.notes.models import Note, Deliberation

from apps.notes.services import mention

from apps.utilisateurs.models import Etudiant

from apps.academique.models import AnneeAcademique, Filiere, Specialite, Classe, Niveau, Matiere

from django.db.models.functions import TruncMonth

from apps.finances.models import Paiement, Depense, Inscription, CaisseSession, TypePaiement, CategorieDepense

from apps.bibliotheque.models import Ressource, Exemplaire, Emprunt, Adherent, Reservation, Penalite, Categorie, StatutExemplaire, StatutReservation

import io

from django.http import HttpResponse

from django.template.loader import render_to_string

from weasyprint import HTML

import openpyxl

from openpyxl.styles import Font, PatternFill









# ---------------------------------------------------------------------------
# Helpers communs
# ---------------------------------------------------------------------------
def _deliberations_filtrees(params, annee_override=None, periode_override=None):
    qs = Deliberation.objects.select_related(
        'etudiant', 'etudiant__specialite__filiere', 'etudiant__specialite',
        'etudiant__classe', 'etudiant__niveau',
    )
    annee_id = annee_override if annee_override is not None else params.get('annee_academique')
    if annee_id:
        qs = qs.filter(annee_academique_id=annee_id)
    periode = periode_override if periode_override is not None else params.get('semestre')
    if periode:
        qs = qs.filter(periode=periode)
    if params.get('filiere'):
        qs = qs.filter(etudiant__specialite__filiere_id=params['filiere'])
    if params.get('specialite'):
        qs = qs.filter(etudiant__specialite_id=params['specialite'])
    if params.get('classe'):
        qs = qs.filter(etudiant__classe_id=params['classe'])
    if params.get('niveau'):
        qs = qs.filter(etudiant__niveau_id=params['niveau'])
    if params.get('sexe'):
        qs = qs.filter(etudiant__sexe=params['sexe'])
    return qs



def _notes_filtrees(params):
    etudiants = Etudiant.objects.filter(statut='ACTIF')
    if params.get('filiere'):
        etudiants = etudiants.filter(specialite__filiere_id=params['filiere'])
    if params.get('specialite'):
        etudiants = etudiants.filter(specialite_id=params['specialite'])
    if params.get('classe'):
        etudiants = etudiants.filter(classe_id=params['classe'])
    if params.get('niveau'):
        etudiants = etudiants.filter(niveau_id=params['niveau'])
    if params.get('sexe'):
        etudiants = etudiants.filter(sexe=params['sexe'])
    notes = Note.objects.filter(etudiant__in=etudiants)
    if params.get('annee_academique'):
        notes = notes.filter(annee_academique_id=params['annee_academique'])
    if params.get('semestre'):
        notes = notes.filter(semestre=params['semestre'])
    if params.get('matiere'):
        notes = notes.filter(matiere_id=params['matiere'])
    return etudiants, notes



def _stats_deliberations(qs):
    """Calcule toutes les statistiques dérivées d'un queryset de Deliberation."""
    total = qs.count()
    if total == 0:
        return {
            'total': 0, 'admis': 0, 'ajournes': 0, 'redoublants': 0,
            'taux_reussite': 0, 'taux_echec': 0, 'taux_rattrapage': 0,
            'moyenne_generale': None, 'mentions': 0, 'taux_mentions': 0,
            'meilleure': None, 'plus_faible': None,
        }
    admis = qs.filter(decision='ADMIS').count()
    ajournes = qs.filter(decision='RATTRAPAGE').count()
    redoublants = qs.filter(decision='REDOUBLANT').count()
    moyennes_valides = qs.exclude(moyenne_generale__isnull=True)
    moyenne_generale = moyennes_valides.aggregate(m=Avg('moyenne_generale'))['m']
    mentions = moyennes_valides.filter(moyenne_generale__gte=10).count()
    meilleure_dl = moyennes_valides.order_by('-moyenne_generale').first()
    plus_faible_dl = moyennes_valides.order_by('moyenne_generale').first()
    def _fiche(dl):
        if not dl:
            return None
        etu = dl.etudiant
        return {
            'nom': f"{etu.nom} {etu.prenom}",
            'matricule': etu.matricule,
            'classe': str(etu.classe) if etu.classe else '—',
            'moyenne': float(dl.moyenne_generale),
        }
    return {
        'total': total, 'admis': admis, 'ajournes': ajournes, 'redoublants': redoublants,
        'taux_reussite': round((admis / total) * 100, 1),
        'taux_echec': round((redoublants / total) * 100, 1),
        'taux_rattrapage': round((ajournes / total) * 100, 1),
        'moyenne_generale': round(float(moyenne_generale), 2) if moyenne_generale is not None else None,
        'mentions': mentions,
        'taux_mentions': round((mentions / moyennes_valides.count()) * 100, 1) if moyennes_valides.count() else 0,
        'meilleure': _fiche(meilleure_dl),
        'plus_faible': _fiche(plus_faible_dl),
    }




def _periode_precedente(annee_id, semestre):
    """Détermine (annee_id, periode) du semestre précédent pour la comparaison."""
    if not semestre or not annee_id:
        return None, None
    try:
        annee = AnneeAcademique.objects.get(pk=annee_id)
    except AnneeAcademique.DoesNotExist:
        return None, None
    if semestre == 'S2':
        return annee_id, 'S1'
    if semestre == 'S1':
        precedente = AnneeAcademique.objects.filter(date_debut__lt=annee.date_debut).order_by('-date_debut').first()
        if precedente:
            return precedente.id, 'S2'
    return None, None




def _delta(actuel, precedent, relatif=False):
    """Renvoie {valeur, sens} ou None si comparaison impossible."""
    if actuel is None or precedent is None:
        return None
    if relatif:
        if precedent == 0:
            return None
        valeur = round(((actuel - precedent) / precedent) * 100, 1)
    else:
        valeur = round(actuel - precedent, 2)
    return {'valeur': valeur, 'sens': 'hausse' if valeur > 0 else ('baisse' if valeur < 0 else 'stable')}




# ---------------------------------------------------------------------------
# Dashboard principal
# ---------------------------------------------------------------------------
@api_view(['GET'])
@permission_classes([IsAuthenticated])
@permission_requise('gerer_notes')
def dashboard_academique(request):
    params = request.GET
    annee_id = params.get('annee_academique')
    if not annee_id:
        annee_active = AnneeAcademique.objects.filter(statut=True).first()
        annee_id = annee_active.id if annee_active else None
    params_effectifs = params.copy()
    params_effectifs['annee_academique'] = annee_id
    # ---- Stats principales (basées sur les délibérations déjà calculées) ----
    deliberations = _deliberations_filtrees(params_effectifs)
    stats = _stats_deliberations(deliberations)
    # ---- Stats de la période précédente (pour comparaison) ----
    annee_prec_id, semestre_prec = _periode_precedente(annee_id, params.get('semestre'))
    stats_prec = None
    if annee_prec_id and semestre_prec:
        deliberations_prec = _deliberations_filtrees(params_effectifs, annee_override=annee_prec_id, periode_override=semestre_prec)
        stats_prec = _stats_deliberations(deliberations_prec)
    comparaisons = {}
    if stats_prec:
        comparaisons = {
            'apprenants_evalues': _delta(stats['total'], stats_prec['total'], relatif=True),
            'taux_reussite': _delta(stats['taux_reussite'], stats_prec['taux_reussite']),
            'taux_echec': _delta(stats['taux_echec'], stats_prec['taux_echec']),
            'moyenne_generale': _delta(stats['moyenne_generale'], stats_prec['moyenne_generale']),
        }
    # ---- Notes brutes (évaluations, tranches, matières) ----
    etudiants, notes_qs = _notes_filtrees(params_effectifs)
    evaluations_realisees = notes_qs.count()
    comparaisons['evaluations_realisees'] = None
    if stats_prec:
        params_prec = params_effectifs.copy()
        params_prec['annee_academique'] = annee_prec_id
        params_prec['semestre'] = semestre_prec
        _, notes_prec_qs = _notes_filtrees(params_prec)
        comparaisons['evaluations_realisees'] = _delta(evaluations_realisees, notes_prec_qs.count(), relatif=True)
    moyennes_deliberations = deliberations.exclude(moyenne_generale__isnull=True).values_list('moyenne_generale', flat=True)
    tranches = {'0-8': 0, '8-10': 0, '10-12': 0, '12-14': 0, '14-16': 0, '16-20': 0}

    for v in moyennes_deliberations:
        v = float(v)
        if v < 8: tranches['0-8'] += 1
        elif v < 10: tranches['8-10'] += 1
        elif v < 12: tranches['10-12'] += 1
        elif v < 14: tranches['12-14'] += 1
        elif v < 16: tranches['14-16'] += 1
        else: tranches['16-20'] += 1
    repartition_tranches = [{'tranche': k, 'nombre': v} for k, v in tranches.items()]
    matieres_concernees = Matiere.objects.filter(id__in=notes_qs.values_list('matiere', flat=True).distinct())
    par_matiere = []

    for mat in matieres_concernees:
        moy = notes_qs.filter(matiere=mat).aggregate(m=Avg('valeur'))['m']
        if moy is not None:
            par_matiere.append({'matiere': mat.nom, 'moyenne': round(float(moy), 2)})
    # ---- Comparaison sexe (sur délibérations) ----
    comparaison_sexe = []
    for code, label in [('M', 'Masculin'), ('F', 'Féminin')]:
        dl_sexe = deliberations.filter(etudiant__sexe=code).exclude(moyenne_generale__isnull=True)
        if dl_sexe.exists():
            comparaison_sexe.append({'sexe': label, 'moyenne': round(float(dl_sexe.aggregate(m=Avg('moyenne_generale'))['m']), 2)})
    # ---- Performance par filière (top 5 + total pour "voir tout") ----
    toutes_filieres = []
    for filiere in Filiere.objects.filter(statut='actif'):
        dl_filiere = deliberations.filter(etudiant__specialite__filiere=filiere)
        s = _stats_deliberations(dl_filiere)
        if s['total'] > 0:
            toutes_filieres.append({'filiere': filiere.nom, 'apprenants': s['total'], 'moyenne': s['moyenne_generale'], 'taux_reussite': s['taux_reussite']})
    toutes_filieres.sort(key=lambda x: x['taux_reussite'], reverse=True)
    # ---- Performance par classe ----
    toutes_classes = []
    for classe in Classe.objects.all():
        dl_classe = deliberations.filter(etudiant__classe=classe)
        s = _stats_deliberations(dl_classe)
        if s['total'] > 0:
            toutes_classes.append({'classe': str(classe), 'apprenants': s['total'], 'moyenne': s['moyenne_generale']})
    toutes_classes.sort(key=lambda x: x['moyenne'], reverse=True)
    # ---- Classement complet des apprenants ----
    classement_complet = []
    for dl in deliberations.exclude(moyenne_generale__isnull=True).order_by('-moyenne_generale'):
        etu = dl.etudiant
        classement_complet.append({
            'nom': f"{etu.nom} {etu.prenom}", 'matricule': etu.matricule,
            'classe': str(etu.classe) if etu.classe else '—',
            'specialite': str(etu.specialite) if etu.specialite else '—',
            'moyenne': float(dl.moyenne_generale), 'mention': mention(dl.moyenne_generale),
            'decision': dl.decision,
        })
    # ---- Alertes ----
    matieres_fort_echec = []
    for mat in matieres_concernees:
        notes_mat = notes_qs.filter(matiere=mat)
        if notes_mat.count() == 0:
            continue
        taux_echec_mat = round((notes_mat.filter(valeur__lt=10).count() / notes_mat.count()) * 100, 1)
        if taux_echec_mat > 30:
            matieres_fort_echec.append({'matiere': mat.nom, 'taux_echec': taux_echec_mat})
    matieres_fort_echec.sort(key=lambda x: x['taux_echec'], reverse=True)
    notes_manquantes = etudiants.count() * matieres_concernees.count() - notes_qs.values('etudiant', 'matiere').distinct().count()
    deliberations_incompletes = deliberations.filter(decision='INCOMPLET').count()
    return Response({
        'kpis': {
            'apprenants_evalues': stats['total'],
            'evaluations_realisees': evaluations_realisees,
            'moyenne_generale': stats['moyenne_generale'],
            'taux_reussite': stats['taux_reussite'],
            'taux_echec': stats['taux_echec'],
            'taux_rattrapage': stats['taux_rattrapage'],
            'redoublants': stats['redoublants'],
            'mentions': stats['mentions'],
            'taux_mentions': stats['taux_mentions'],
            'meilleure_moyenne': stats['meilleure'],
            'plus_faible_moyenne': stats['plus_faible'],
        },
        'comparaisons': comparaisons,
        'repartition_resultats': [
            {'statut': 'Admis', 'nb': stats['admis']},
            {'statut': 'Ajournés', 'nb': stats['ajournes']},
            {'statut': 'Redoublants', 'nb': stats['redoublants']},
        ],
        'repartition_tranches': repartition_tranches,
        'moyenne_par_matiere': par_matiere,
        'comparaison_sexe': comparaison_sexe,
        'top_filieres': toutes_filieres[:5],
        'toutes_filieres': toutes_filieres,
        'top_classes': toutes_classes[:5],
        'toutes_classes': toutes_classes,
        'top_apprenants': classement_complet[:10],
        'classement_complet': classement_complet,
        'alertes': {
            'apprenants_en_difficulte': deliberations.filter(decision='REDOUBLANT').count(),
            'matieres_fort_taux_echec': matieres_fort_echec[:10],
            'notes_manquantes': max(notes_manquantes, 0),
            'deliberations_incompletes': deliberations_incompletes,
        },
    })




@api_view(['GET'])
@permission_classes([IsAuthenticated])
@permission_requise('gerer_notes')
def filtres_academique(request):
    return Response({
        'annees_academiques': list(AnneeAcademique.objects.values('id', 'libelle')),
        'filieres': list(Filiere.objects.filter(statut='actif').values('id', 'nom')),
        'specialites': list(Specialite.objects.filter(statut='actif').values('id', 'nom', 'filiere_id')),
        'classes': [{'id': c.id, 'nom': str(c)} for c in Classe.objects.all()],
        'niveaux': list(Niveau.objects.values('id', 'nom')),
        'matieres': list(Matiere.objects.filter(statut='actif').values('id', 'nom')),
    })








# ---------------------------------------------------------------------------
# Exports
# ---------------------------------------------------------------------------
@api_view(['GET'])
@permission_classes([IsAuthenticated])
@permission_requise('gerer_notes')
def export_academique_pdf(request):
    reponse_dashboard = dashboard_academique(request._request)
    donnees = reponse_dashboard.data
    html_string = render_to_string('statistiques/academique_pdf.html', {
        'donnees': donnees,
        'date_generation': timezone.now().strftime('%d/%m/%Y à %H:%M'),
    })
    pdf_file = HTML(string=html_string, base_url=request.build_absolute_uri('/')).write_pdf()
    response = HttpResponse(pdf_file, content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="dashboard_academique.pdf"'
    return response


@api_view(['GET'])
@permission_classes([IsAuthenticated])
@permission_requise('gerer_notes')
def export_academique_excel(request):
    reponse_dashboard = dashboard_academique(request._request)
    donnees = reponse_dashboard.data
    wb = openpyxl.Workbook()
    entete_style = Font(bold=True, color='FFFFFF')
    entete_fill = PatternFill(start_color='400C7C', end_color='400C7C', fill_type='solid')
    ws1 = wb.active
    ws1.title = 'KPI'
    ws1.append(['Indicateur', 'Valeur'])
    for cell in ws1[1]:
        cell.font = entete_style
        cell.fill = entete_fill
    for cle, valeur in donnees['kpis'].items():
        if isinstance(valeur, dict):
            valeur = f"{valeur.get('nom', '')} ({valeur.get('moyenne', '')})" if valeur else '—'
        ws1.append([cle, valeur])
    ws2 = wb.create_sheet('Classement complet')
    ws2.append(['Nom', 'Matricule', 'Classe', 'Spécialité', 'Moyenne', 'Mention', 'Décision'])
    for cell in ws2[1]:
        cell.font = entete_style
        cell.fill = entete_fill
    for e in donnees['classement_complet']:
        ws2.append([e['nom'], e['matricule'], e['classe'], e['specialite'], e['moyenne'], e['mention'], e['decision']])
    ws3 = wb.create_sheet('Filières')
    ws3.append(['Filière', 'Apprenants', 'Moyenne', 'Taux de réussite'])
    for cell in ws3[1]:
        cell.font = entete_style
        cell.fill = entete_fill
    for f in donnees['toutes_filieres']:
        ws3.append([f['filiere'], f['apprenants'], f['moyenne'], f['taux_reussite']])
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(buffer.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="dashboard_academique.xlsx"'
    return response























def _appliquer_filtres_emprunts(emprunts, params):
    categorie = params.get('categorie')
    if categorie:
        emprunts = emprunts.filter(exemplaire__ressource__categorie_id=categorie)
    type_utilisateur = params.get('type_utilisateur')
    if type_utilisateur == 'ETUDIANT':
        emprunts = emprunts.filter(adherent__etudiant__isnull=False)
    elif type_utilisateur == 'FORMATEUR':
        emprunts = emprunts.filter(adherent__formateur__isnull=False)
    elif type_utilisateur == 'PERSONNEL':
        emprunts = emprunts.filter(adherent__personnel__isnull=False, adherent__formateur__isnull=True)
    statut = params.get('statut_emprunt')
    if statut:
        emprunts = emprunts.filter(statut=statut)
    filiere = params.get('filiere')
    if filiere:
        emprunts = emprunts.filter(
            Q(adherent__etudiant__specialite__filiere_id=filiere) |
            Q(adherent__formateur__personnel__isnull=False)  # formateurs non filtrables par filière d'étudiant, ignoré silencieusement
        )
    specialite = params.get('specialite')
    if specialite:
        emprunts = emprunts.filter(adherent__etudiant__specialite_id=specialite)
    date_debut = params.get('date_debut')
    if date_debut:
        emprunts = emprunts.filter(date_emprunt__gte=date_debut)
    date_fin = params.get('date_fin')
    if date_fin:
        emprunts = emprunts.filter(date_emprunt__lte=date_fin)
    return emprunts



@api_view(['GET'])
@permission_classes([IsAuthenticated])
@permission_requise('gerer_bibliotheque_ressources')
def dashboard_bibliotheque(request):
    emprunts = Emprunt.objects.select_related('adherent', 'exemplaire__ressource__categorie').all()
    emprunts = _appliquer_filtres_emprunts(emprunts, request.GET)
    # ---- KPI ----
    kpis = {
        'nb_livres': Ressource.objects.filter(retiree=False).count(),
        'nb_exemplaires': Exemplaire.objects.count(),
        'exemplaires_disponibles': Exemplaire.objects.filter(statut=StatutExemplaire.DISPONIBLE).count(),
        'exemplaires_empruntes': Exemplaire.objects.filter(statut=StatutExemplaire.EMPRUNTE).count(),
        'emprunts_en_cours': emprunts.filter(statut__in=['EN_COURS', 'EN_RETARD']).count(),
        'emprunts_en_retard': sum(1 for e in emprunts.filter(statut__in=['EN_COURS', 'EN_RETARD']) if e.est_en_retard),
        'lecteurs_actifs': Adherent.objects.filter(statut='ACTIF', emprunts__in=emprunts).distinct().count(),
        'emprunts_periode': emprunts.count(),
        'reservations_attente': Reservation.objects.filter(statut=StatutReservation.EN_ATTENTE).count(),
        'penalites_generees': Penalite.objects.aggregate(t=Count('id'))['t'],
    }
    # ---- Évolution des emprunts dans le temps (par mois) ----
    from django.db.models.functions import TruncMonth
    evolution = (
        emprunts.annotate(mois=TruncMonth('date_emprunt'))
        .values('mois').annotate(nb=Count('id')).order_by('mois')
    )
    evolution_data = [{'mois': e['mois'].strftime('%b %Y'), 'nb': e['nb']} for e in evolution if e['mois']]
    # ---- Emprunts par catégorie ----
    par_categorie = (
        emprunts.values('exemplaire__ressource__categorie__nom')
        .annotate(nb=Count('id')).order_by('-nb')
    )
    par_categorie_data = [{'categorie': c['exemplaire__ressource__categorie__nom'] or 'Non classé', 'nb': c['nb']} for c in par_categorie]
    # ---- Top livres empruntés ----
    top_livres = (
        emprunts.values('exemplaire__ressource__titre')
        .annotate(nb=Count('id')).order_by('-nb')[:10]
    )
    top_livres_data = [{'titre': t['exemplaire__ressource__titre'], 'nb': t['nb']} for t in top_livres]
    # ---- Répartition des exemplaires (donut) ----
    repartition_exemplaires = [
        {'statut': 'Disponibles', 'nb': Exemplaire.objects.filter(statut=StatutExemplaire.DISPONIBLE).count()},
        {'statut': 'Empruntés', 'nb': Exemplaire.objects.filter(statut=StatutExemplaire.EMPRUNTE).count()},
        {'statut': 'Réservés', 'nb': Exemplaire.objects.filter(statut=StatutExemplaire.RESERVE).count()},
        {'statut': 'Perdus', 'nb': Exemplaire.objects.filter(statut=StatutExemplaire.PERDU).count()},
        {'statut': 'Endommagés', 'nb': Exemplaire.objects.filter(statut__in=[StatutExemplaire.ENDOMMAGE, StatutExemplaire.EN_REPARATION]).count()},
    ]
    # ---- Emprunts par type d'utilisateur ----
    par_type_utilisateur = [
        {'type': 'Apprenants', 'nb': emprunts.filter(adherent__etudiant__isnull=False).count()},
        {'type': 'Formateurs', 'nb': emprunts.filter(adherent__formateur__isnull=False).count()},
        {'type': 'Personnels', 'nb': emprunts.filter(adherent__personnel__isnull=False, adherent__formateur__isnull=True).count()},
    ]
    # ---- Tableau retards ----
    retards = [e for e in emprunts.filter(statut__in=['EN_COURS', 'EN_RETARD']) if e.est_en_retard]
    retards.sort(key=lambda e: e.jours_de_retard, reverse=True)
    retards_data = [{
        'utilisateur': str(e.adherent.personne), 'livre': e.exemplaire.ressource.titre, 'jours_retard': e.jours_de_retard,
    } for e in retards[:15]]
    # ---- Tableau réservations en attente ----
    reservations_attente = Reservation.objects.filter(statut=StatutReservation.EN_ATTENTE).select_related('adherent', 'ressource').order_by('date_reservation')[:15]
    reservations_data = [{
        'utilisateur': str(r.adherent.personne), 'ressource': r.ressource.titre,
        'date': r.date_reservation.strftime('%d/%m/%Y'), 'position': r.position_file,
    } for r in reservations_attente]
    # ---- Alertes ----
    exemplaires_perdus = Exemplaire.objects.filter(statut=StatutExemplaire.PERDU).count()
    exemplaires_endommages = Exemplaire.objects.filter(statut=StatutExemplaire.ENDOMMAGE).count()
    penalites_impayees = Penalite.objects.filter(payee=False).count()
    livres_tres_demandes = sum(1 for t in top_livres_data if t['nb'] >= 5)
    alertes = {
        'nb_retards': kpis['emprunts_en_retard'],
        'livres_tres_demandes': livres_tres_demandes,
        'exemplaires_perdus': exemplaires_perdus,
        'exemplaires_endommages': exemplaires_endommages,
        'penalites_impayees': penalites_impayees,
    }
    return Response({
        'kpis': kpis,
        'evolution_emprunts': evolution_data,
        'emprunts_par_categorie': par_categorie_data,
        'top_livres': top_livres_data,
        'repartition_exemplaires': repartition_exemplaires,
        'emprunts_par_type_utilisateur': par_type_utilisateur,
        'retards': retards_data,
        'reservations_attente': reservations_data,
        'alertes': alertes,
    })



@api_view(['GET'])
@permission_classes([IsAuthenticated])
@permission_requise('gerer_bibliotheque_ressources')
def filtres_bibliotheque(request):
    return Response({
        'annees_academiques': list(AnneeAcademique.objects.values('id', 'libelle')),
        'filieres': list(Filiere.objects.filter(statut='actif').values('id', 'nom')),
        'specialites': list(Specialite.objects.filter(statut='actif').values('id', 'nom')),
        'categories': list(Categorie.objects.values('id', 'nom')),
    })
























from apps.documents.models import Document, Diplome, Certificat, TypeCertificat
def _appliquer_filtres_documents(params):
    date_debut = params.get('date_debut')
    date_fin = params.get('date_fin')
    categorie = params.get('categorie')
    concerne = params.get('concerne')  # 'etudiant' | 'personnel'
    responsable = params.get('responsable')
    documents = Document.objects.select_related('concerne_etudiant', 'concerne_personnel', 'ajoute_par')
    if date_debut:
        documents = documents.filter(date_ajout__gte=date_debut)
    if date_fin:
        documents = documents.filter(date_ajout__lte=date_fin)
    if categorie:
        documents = documents.filter(categorie=categorie)
    if concerne == 'etudiant':
        documents = documents.filter(concerne_etudiant__isnull=False)
    elif concerne == 'personnel':
        documents = documents.filter(concerne_personnel__isnull=False)
    if responsable:
        documents = documents.filter(ajoute_par_id=responsable)
    return documents
@api_view(['GET'])
@permission_classes([IsAuthenticated])
@permission_requise('gerer_documents')
def dashboard_documents(request):
    params = request.GET
    documents = _appliquer_filtres_documents(params)
    date_debut = params.get('date_debut')
    date_fin = params.get('date_fin')
    diplomes = Diplome.objects.all()
    certificats = Certificat.objects.all()
    if date_debut:
        diplomes = diplomes.filter(date_creation__gte=date_debut)
        certificats = certificats.filter(date_creation__gte=date_debut)
    if date_fin:
        diplomes = diplomes.filter(date_creation__lte=date_fin)
        certificats = certificats.filter(date_creation__lte=date_fin)
    # ---- KPI ----
    total_documents = documents.count() + diplomes.count() + certificats.count()
    debut_periode = timezone.now() - timezone.timedelta(days=30)
    documents_periode = (
        documents.filter(date_ajout__gte=debut_periode).count() +
        diplomes.filter(date_creation__gte=debut_periode).count() +
        certificats.filter(date_creation__gte=debut_periode).count()
    )
    kpis = {
        'total_documents': total_documents,
        'documents_ajoutes_periode': documents_periode,
        'documents_divers': documents.count(),
        'diplomes_delivres': diplomes.count(),
        'diplomes_valides': diplomes.filter(statut='valide').count(),
        'diplomes_revoques': diplomes.filter(statut='revoque').count(),
        'certificats_delivres': certificats.count(),
        'documents_lies_etudiants': documents.filter(concerne_etudiant__isnull=False).count(),
        'documents_lies_personnel': documents.filter(concerne_personnel__isnull=False).count(),
    }
    # ---- Évolution du volume dans le temps ----
    from django.db.models.functions import TruncMonth
    evol_docs = documents.annotate(mois=TruncMonth('date_ajout')).values('mois').annotate(nb=Count('id'))
    evol_dip = diplomes.annotate(mois=TruncMonth('date_creation')).values('mois').annotate(nb=Count('id'))
    evol_cert = certificats.annotate(mois=TruncMonth('date_creation')).values('mois').annotate(nb=Count('id'))
    mois_map = {}
    for e in evol_docs:
        if e['mois']:
            mois_map[e['mois']] = mois_map.get(e['mois'], 0) + e['nb']
    for e in evol_dip:
        if e['mois']:
            mois_map[e['mois']] = mois_map.get(e['mois'], 0) + e['nb']
    for e in evol_cert:
        if e['mois']:
            mois_map[e['mois']] = mois_map.get(e['mois'], 0) + e['nb']
    evolution = [{'mois': m.strftime('%b %Y'), 'nb': n} for m, n in sorted(mois_map.items())]
    # ---- Documents par catégorie ----
    par_categorie = documents.exclude(categorie__isnull=True).exclude(categorie='').values('categorie').annotate(nb=Count('id')).order_by('-nb')
    par_categorie_data = [{'categorie': c['categorie'], 'nb': c['nb']} for c in par_categorie]
    # ---- Répartition par type de document ----
    repartition_type = [
        {'type': 'Documents divers', 'nb': documents.count()},
        {'type': 'Diplômes', 'nb': diplomes.count()},
        {'type': 'Certificats', 'nb': certificats.count()},
    ]
    # ---- Par responsable (qui ajoute le plus de documents) ----
    par_responsable = (
        documents.exclude(ajoute_par__isnull=True).values('ajoute_par__username')
        .annotate(nb=Count('id')).order_by('-nb')[:8]
    )
    par_responsable_data = [{'utilisateur': r['ajoute_par__username'], 'nb': r['nb']} for r in par_responsable]
    # ---- Répartition des diplômes par statut ----
    repartition_diplomes = [
        {'statut': 'Valides', 'nb': diplomes.filter(statut='valide').count()},
        {'statut': 'Révoqués', 'nb': diplomes.filter(statut='revoque').count()},
    ]
    # ---- Répartition certificats par type ----
    par_type_certificat = (
        certificats.values('type_certificat__nom').annotate(nb=Count('id')).order_by('-nb')
    )
    par_type_certificat_data = [{'type': t['type_certificat__nom'], 'nb': t['nb']} for t in par_type_certificat]
    # ---- Activité récente ----
    activites = []
    for d in documents.order_by('-date_ajout')[:5]:
        activites.append({'texte': f"Document ajouté : {d.titre}", 'date': d.date_ajout.isoformat(), 'type': 'document'})
    for d in diplomes.order_by('-date_creation')[:5]:
        activites.append({'texte': f"Diplôme délivré : {d.etudiant}", 'date': d.date_creation.isoformat(), 'type': 'diplome'})
    for c in certificats.order_by('-date_creation')[:5]:
        activites.append({'texte': f"Certificat généré : {c.type_certificat} — {c.etudiant}", 'date': c.date_creation.isoformat(), 'type': 'certificat'})
    activites.sort(key=lambda a: a['date'], reverse=True)
    activites = activites[:10]
    # ---- Alertes (adaptées : uniquement ce qui est réellement calculable) ----
    diplomes_revoques_recents = diplomes.filter(statut='revoque', date_creation__gte=debut_periode).count()
    documents_sans_categorie = documents.filter(Q(categorie__isnull=True) | Q(categorie='')).count()
    alertes = {
        'diplomes_revoques_recents': diplomes_revoques_recents,
        'documents_sans_categorie': documents_sans_categorie,
    }
    return Response({
        'kpis': kpis,
        'evolution': evolution,
        'documents_par_categorie': par_categorie_data,
        'repartition_par_type': repartition_type,
        'documents_par_responsable': par_responsable_data,
        'repartition_diplomes': repartition_diplomes,
        'certificats_par_type': par_type_certificat_data,
        'activite_recente': activites,
        'alertes': alertes,
    })
@api_view(['GET'])
@permission_classes([IsAuthenticated])
@permission_requise('gerer_documents')
def filtres_documents(request):
    from apps.authentification.models import User
    categories = Document.objects.exclude(categorie__isnull=True).exclude(categorie='').values_list('categorie', flat=True).distinct()
    responsables_ids = Document.objects.exclude(ajoute_par__isnull=True).values_list('ajoute_par', flat=True).distinct()
    return Response({
        'categories': list(categories),
        'responsables': list(User.objects.filter(id__in=responsables_ids).values('id', 'username')),
    })























def _appliquer_filtres_finance(qs, params, champ_date):
    date_debut = params.get('date_debut')
    if date_debut:
        qs = qs.filter(**{f'{champ_date}__gte': date_debut})
    date_fin = params.get('date_fin')
    if date_fin:
        qs = qs.filter(**{f'{champ_date}__lte': date_fin})
    return qs


@api_view(['GET'])
@permission_classes([IsAuthenticated])
@permission_requise('gerer_finances')
def dashboard_finance(request):
    params = request.GET
    paiements = Paiement.objects.select_related('inscription__etudiant', 'inscription__classe__filiere', 'type_paiement', 'caisse_session')
    paiements = _appliquer_filtres_finance(paiements, params, 'date_paiement')
    type_frais = params.get('type_frais')
    if type_frais:
        paiements = paiements.filter(type_paiement_id=type_frais)
    mode = params.get('mode_paiement')
    if mode:
        paiements = paiements.filter(mode_paiement=mode)
    statut_paiement = params.get('statut_paiement')
    if statut_paiement:
        paiements = paiements.filter(statut=statut_paiement)
    caisse = params.get('caisse')
    if caisse:
        paiements = paiements.filter(caisse_session_id=caisse)
    caissier = params.get('caissier')
    if caissier:
        paiements = paiements.filter(enregistre_par_id=caissier)
    filiere = params.get('filiere')
    if filiere:
        paiements = paiements.filter(inscription__classe__filiere_id=filiere)
    specialite = params.get('specialite')
    if specialite:
        paiements = paiements.filter(inscription__classe__specialite_id=specialite)
    annee_academique = params.get('annee_academique')
    if annee_academique:
        paiements = paiements.filter(inscription__annee_academique_id=annee_academique)
    depenses = Depense.objects.select_related('categorie', 'caisse_session')
    depenses = _appliquer_filtres_finance(depenses, params, 'date_depense')
    if caisse:
        depenses = depenses.filter(caisse_session_id=caisse)
    # ---- KPI ----
    total_revenus = paiements.filter(statut='VALIDE').aggregate(t=Sum('montant'))['t'] or Decimal('0')
    total_depenses = depenses.exclude(statut='REJETEE').aggregate(t=Sum('montant'))['t'] or Decimal('0')
    total_remboursements = paiements.filter(statut='REMBOURSE').aggregate(t=Sum('montant'))['t'] or Decimal('0')
    inscriptions = Inscription.objects.exclude(statut='ANNULEE')
    if annee_academique:
        inscriptions = inscriptions.filter(annee_academique_id=annee_academique)
    total_impayes = sum((i.reste_a_payer for i in inscriptions), Decimal('0'))
    nb_paiements_attente = sum(1 for i in inscriptions if i.statut_paiement != 'PAYE')
    session_ouverte = CaisseSession.objects.filter(statut='OUVERTE').first()
    solde_caisse_actuel = session_ouverte.solde_theorique if session_ouverte else Decimal('0')
    kpis = {
        'total_revenus': str(total_revenus),
        'total_depenses': str(total_depenses),
        'solde_financier': str(total_revenus - total_depenses),
        'excedent_deficit': str(total_revenus - total_depenses),
        'total_paiements': paiements.filter(statut='VALIDE').count(),
        'total_impayes': str(total_impayes),
        'total_remboursements': str(total_remboursements),
        'solde_caisse_actuel': str(solde_caisse_actuel),
        'nb_paiements_en_attente': nb_paiements_attente,
    }
    # ---- Évolution revenus/dépenses ----
    revenus_mensuels = (
        paiements.filter(statut='VALIDE').annotate(mois=TruncMonth('date_paiement'))
        .values('mois').annotate(total=Sum('montant')).order_by('mois')
    )
    depenses_mensuelles = (
        depenses.exclude(statut='REJETEE').annotate(mois=TruncMonth('date_depense'))
        .values('mois').annotate(total=Sum('montant')).order_by('mois')
    )
    mois_set = sorted(set(
        [r['mois'] for r in revenus_mensuels] + [d['mois'] for d in depenses_mensuelles]
    ))
    rev_map = {r['mois']: float(r['total']) for r in revenus_mensuels}
    dep_map = {d['mois']: float(d['total']) for d in depenses_mensuelles}
    evolution = [{
        'mois': m.strftime('%b %Y'),
        'revenus': round(rev_map.get(m, 0), 2),
        'depenses': round(dep_map.get(m, 0), 2),
    } for m in mois_set]
    # ---- Répartition revenus par type de frais ----
    par_type_frais = (
        paiements.filter(statut='VALIDE').values('type_paiement__nom')
        .annotate(total=Sum('montant')).order_by('-total')
    )
    par_type_frais_data = [{'type': t['type_paiement__nom'], 'montant': float(t['total'])} for t in par_type_frais]
    # ---- Répartition dépenses par catégorie ----
    par_categorie = (
        depenses.exclude(statut='REJETEE').values('categorie__nom')
        .annotate(total=Sum('montant')).order_by('-total')
    )
    par_categorie_data = [{'categorie': c['categorie__nom'], 'montant': float(c['total'])} for c in par_categorie]
    # ---- Revenus par filière ----
    par_filiere = (
        paiements.filter(statut='VALIDE').values('inscription__classe__filiere__nom')
        .annotate(total=Sum('montant')).order_by('-total')
    )
    par_filiere_data = [{'filiere': f['inscription__classe__filiere__nom'] or 'Non défini', 'montant': float(f['total'])} for f in par_filiere]
    # ---- Répartition des paiements par statut ----
    repartition_statuts_paiement = [
        {'statut': 'Validé', 'nb': paiements.filter(statut='VALIDE').count()},
        {'statut': 'Annulé', 'nb': paiements.filter(statut='ANNULE').count()},
        {'statut': 'Remboursé', 'nb': paiements.filter(statut='REMBOURSE').count()},
    ]
    repartition_inscriptions = [
        {'statut': 'Payé', 'nb': sum(1 for i in inscriptions if i.statut_paiement == 'PAYE')},
        {'statut': 'Partiel', 'nb': sum(1 for i in inscriptions if i.statut_paiement == 'PARTIEL')},
        {'statut': 'Non payé', 'nb': sum(1 for i in inscriptions if i.statut_paiement == 'NON_PAYE')},
    ]
    # ---- Évolution des impayés (approximation par mois d'inscription) ----
    impayes_par_mois = {}
    for i in inscriptions:
        if i.reste_a_payer > 0:
            mois_key = i.date_inscription.strftime('%b %Y')
            impayes_par_mois[mois_key] = impayes_par_mois.get(mois_key, 0) + float(i.reste_a_payer)
    evolution_impayes = [{'mois': k, 'montant': round(v, 2)} for k, v in impayes_par_mois.items()]
    # ---- Transactions récentes ----
    transactions = []
    for p in paiements.order_by('-date_paiement')[:8]:
        transactions.append({
            'type': 'Paiement', 'libelle': f"{p.numero_recu} — {p.inscription.etudiant}",
            'montant': str(p.montant), 'date': p.date_paiement.isoformat(), 'statut': p.statut,
        })
    for d in depenses.order_by('-date_depense')[:8]:
        transactions.append({
            'type': 'Dépense', 'libelle': d.libelle,
            'montant': str(-d.montant), 'date': d.date_depense.isoformat(), 'statut': d.statut,
        })
    transactions.sort(key=lambda t: t['date'], reverse=True)
    transactions = transactions[:12]
    # ---- Situation de caisse ----
    derniere_session = session_ouverte or CaisseSession.objects.order_by('-date_session').first()
    situation_caisse = None
    if derniere_session:
        situation_caisse = {
            'statut': derniere_session.statut,
            'solde_theorique': str(derniere_session.solde_theorique),
            'solde_reel': str(derniere_session.solde_reel_fermeture) if derniere_session.solde_reel_fermeture is not None else None,
            'ecart': str(derniere_session.ecart) if derniere_session.ecart is not None else None,
            'total_entrees': str(derniere_session.total_encaisse_especes),
            'total_sorties': str(derniere_session.total_depense_especes),
            'date_session': derniere_session.date_session.isoformat(),
            'heure_ouverture': derniere_session.heure_ouverture.strftime('%H:%M') if derniere_session.heure_ouverture else None,
            'heure_fermeture': derniere_session.heure_fermeture.strftime('%H:%M') if derniere_session.heure_fermeture else None,
            'ouverte_par': derniere_session.ouverte_par.username if derniere_session.ouverte_par else None,
            'fermee_par': derniere_session.fermee_par.username if derniere_session.fermee_par else None,
        }
    # ---- Alertes ----
    impayes_importants = sum(1 for i in inscriptions if i.reste_a_payer > 100000)
    ecarts_caisse = CaisseSession.objects.filter(statut='FERMEE').exclude(solde_reel_fermeture=None)
    nb_ecarts = sum(1 for s in ecarts_caisse if s.ecart and s.ecart != 0)
    caisses_non_fermees = CaisseSession.objects.filter(statut='OUVERTE', date_session__lt=timezone.localdate()).count()
    depenses_importantes = depenses.filter(montant__gt=100000).count()
    alertes = {
        'impayes_importants': impayes_importants,
        'ecarts_caisse': nb_ecarts,
        'caisses_non_fermees': caisses_non_fermees,
        'remboursements': paiements.filter(statut='REMBOURSE').count(),
        'depenses_importantes': depenses_importantes,
    }
    return Response({
        'kpis': kpis,
        'evolution': evolution,
        'revenus_par_type_frais': par_type_frais_data,
        'depenses_par_categorie': par_categorie_data,
        'revenus_par_filiere': par_filiere_data,
        'repartition_statuts_paiement': repartition_statuts_paiement,
        'repartition_inscriptions': repartition_inscriptions,
        'evolution_impayes': evolution_impayes,
        'transactions_recentes': transactions,
        'situation_caisse': situation_caisse,
        'alertes': alertes,
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
@permission_requise('gerer_finances')
def filtres_finance(request):
    from apps.authentification.models import User
    caissiers_ids = Paiement.objects.exclude(enregistre_par__isnull=True).values_list('enregistre_par', flat=True).distinct()
    return Response({
        'annees_academiques': list(AnneeAcademique.objects.values('id', 'libelle')),
        'filieres': list(Filiere.objects.filter(statut='actif').values('id', 'nom')),
        'specialites': list(Specialite.objects.filter(statut='actif').values('id', 'nom')),
        'types_paiement': list(TypePaiement.objects.values('id', 'nom')),
        'categories_depense': list(CategorieDepense.objects.values('id', 'nom')),
        'caisses': [{'id': c.id, 'libelle': f"Caisse du {c.date_session}"} for c in CaisseSession.objects.order_by('-date_session')[:30]],
        'caissiers': list(User.objects.filter(id__in=caissiers_ids).values('id', 'username')),
    })


